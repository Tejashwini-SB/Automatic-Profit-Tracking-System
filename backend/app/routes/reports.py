import os
import tempfile
import openpyxl
from fastapi import APIRouter, Depends
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from .. import database, auth, models

router = APIRouter(
    prefix="/api/reports",
    tags=['Reports']
)

def get_db():
    db = database.SessionLocal()
    try:
        yield db
    finally:
        db.close()

def create_excel_response(filename: str, sheets_data: dict):
    """
    sheets_data: { "Sheet Title": [ [row1_val1, row1_val2], [row2_val1, row2_val2] ] }
    """
    wb = openpyxl.Workbook()
    # remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    for title, rows in sheets_data.items():
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
            
        # Basic Styling: Bold headers
        for col in range(1, len(rows[0]) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid") # Indigo header
            
    temp_fd, temp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(temp_fd)
    
    wb.save(temp_path)
    wb.close()
    
    return FileResponse(
        temp_path,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        filename=filename
    )

@router.get("/inventory")
def export_inventory(db: Session = Depends(get_db), current_user = Depends(auth.get_current_user)):
    products = db.query(models.Product).order_by(models.Product.product_name.asc()).all()
    rows = [
        ["Product ID", "Product Name", "Category", "Cost Price", "Selling Price", "Quantity In Stock", "Inventory Valuation"]
    ]
    for p in products:
        rows.append([
            p.id, p.product_name, p.category, p.cost_price, p.selling_price, p.quantity,
            p.cost_price * p.quantity
        ])
        
    return create_excel_response("Inventory_Report.xlsx", {"Inventory Valuation": rows})

@router.get("/purchases")
def export_purchases(db: Session = Depends(get_db), current_user = Depends(auth.get_current_user)):
    purchases = db.query(models.Purchase).order_by(models.Purchase.purchased_at.desc()).all()
    rows = [
        ["Purchase ID", "Product Name", "Supplier", "Quantity Purchased", "Purchase Cost per Unit", "Total Expense", "Purchased At"]
    ]
    for p in purchases:
        p_name = p.product.product_name if p.product else "Unknown"
        rows.append([
            p.id, p_name, p.supplier, p.quantity, p.purchase_price, p.total_amount,
            p.purchased_at.strftime("%Y-%m-%d %H:%M:%S")
        ])
        
    return create_excel_response("Purchases_Report.xlsx", {"Purchase Logs": rows})

@router.get("/sales")
def export_sales(db: Session = Depends(get_db), current_user = Depends(auth.get_current_user)):
    sales = db.query(models.Sale).order_by(models.Sale.sold_at.desc()).all()
    rows = [
        ["Sale ID", "Product Name", "Quantity Sold", "Selling Price per Unit", "Total Sale Revenue", "Net Profit Margin", "Sold At"]
    ]
    for s in sales:
        p_name = s.product.product_name if s.product else "Unknown"
        rows.append([
            s.id, p_name, s.quantity, s.selling_price, s.total_sale, s.profit,
            s.sold_at.strftime("%Y-%m-%d %H:%M:%S")
        ])
        
    return create_excel_response("Sales_Report.xlsx", {"Sales Logs": rows})

@router.get("/profit")
def export_profit(db: Session = Depends(get_db), current_user = Depends(auth.get_current_user)):
    summaries = db.query(models.ProfitSummary).order_by(models.ProfitSummary.date.desc()).all()
    rows = [
        ["Summary ID", "Date", "Total Revenue", "Total Expense", "Net Profit"]
    ]
    for s in summaries:
        rows.append([
            s.id, s.date.strftime("%Y-%m-%d"), s.revenue, s.expense, s.profit
        ])
        
    return create_excel_response("Daily_Profit_Report.xlsx", {"Profit History": rows})
